import subprocess
import tempfile
import unittest
import zipfile
from pathlib import Path
from shutil import copyfile
from unittest import mock

from docx import Document
from docx.shared import Inches
from PIL import Image

from ppt_watermark.config import WatermarkConfig
from ppt_watermark.word_processor import WordWatermarkProcessor


class WordProcessorTests(unittest.TestCase):
    def test_process_file_can_fallback_to_office_normalized_docx(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            root = Path(tmp_dir)
            source = root / "invalid.docx"
            source.write_bytes(b"not-a-zip-docx")
            output = root / "invalid_watermarked.docx"

            image_path = root / "sample.png"
            Image.new("RGB", (160, 90), color=(255, 255, 255)).save(image_path)
            normalized_source = root / "normalized_source.docx"
            doc = Document()
            doc.add_paragraph("fallback test")
            doc.add_picture(str(image_path), width=Inches(2))
            doc.save(str(normalized_source))

            def fake_normalize(input_path: Path, target_path: Path) -> str:
                self.assertEqual(input_path, source.resolve())
                copyfile(normalized_source, target_path)
                return "OK|Word|Word.Application"

            processor = WordWatermarkProcessor()
            with mock.patch.object(
                processor,
                "_normalize_with_local_office",
                side_effect=fake_normalize,
            ) as normalize:
                stats = processor.process_file(source, output, WatermarkConfig(text="测试"))

            normalize.assert_called_once()
            self.assertTrue(output.exists())
            self.assertEqual(stats.total_images, 1)
            self.assertEqual(stats.watermarked_images, 1)

    def test_process_file_surfaces_office_fallback_failure_clearly(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            root = Path(tmp_dir)
            source = root / "invalid.docx"
            output = root / "invalid_watermarked.docx"
            source.write_bytes(b"not-a-zip-docx")

            processor = WordWatermarkProcessor()
            with mock.patch.object(
                processor,
                "_normalize_with_local_office",
                side_effect=ValueError("自动调用本机 Word/WPS 另存失败"),
            ):
                with self.assertRaisesRegex(ValueError, "自动调用本机 Word/WPS 另存失败"):
                    processor.process_file(source, output, WatermarkConfig(text="测试"))

    def test_process_file_supports_standard_docx_with_images(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            root = Path(tmp_dir)
            image_path = root / "sample.png"
            Image.new("RGB", (160, 90), color=(255, 255, 255)).save(image_path)

            source = root / "source.docx"
            output = root / "source_watermarked.docx"
            doc = Document()
            doc.add_paragraph("watermark test")
            doc.add_picture(str(image_path), width=Inches(2))
            doc.save(str(source))

            processor = WordWatermarkProcessor(min_area_pct=0)
            stats = processor.process_file(source, output, WatermarkConfig(text="测试水印"))

            self.assertTrue(output.exists())
            self.assertEqual(stats.total_images, 1)
            self.assertEqual(stats.watermarked_images, 1)

    def test_process_file_includes_header_images(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            root = Path(tmp_dir)
            image_path = root / "header.png"
            Image.new("RGB", (160, 90), color=(255, 255, 255)).save(image_path)

            source = root / "header_source.docx"
            output = root / "header_source_watermarked.docx"
            doc = Document()
            doc.add_paragraph("header image test")
            header = doc.sections[0].header
            run = header.paragraphs[0].add_run()
            run.add_picture(str(image_path), width=Inches(1.5))
            doc.save(str(source))

            processor = WordWatermarkProcessor(min_area_pct=0)
            stats = processor.process_file(source, output, WatermarkConfig(text="测试水印"))

            self.assertTrue(output.exists())
            self.assertEqual(stats.total_images, 1)
            self.assertEqual(stats.watermarked_images, 1)

    def test_same_source_image_with_different_display_sizes_creates_distinct_media(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            root = Path(tmp_dir)
            image_path = root / "shared.png"
            Image.new("RGB", (600, 300), color=(255, 255, 255)).save(image_path)

            source = root / "shared.docx"
            output = root / "shared_watermarked.docx"
            doc = Document()
            doc.add_picture(str(image_path), width=Inches(2))
            doc.add_picture(str(image_path), width=Inches(4))
            doc.save(str(source))

            processor = WordWatermarkProcessor(min_area_pct=0)
            stats = processor.process_file(source, output, WatermarkConfig(text="测试水印"))

            self.assertEqual(stats.total_images, 2)
            self.assertEqual(stats.watermarked_images, 2)

            with zipfile.ZipFile(output) as package:
                media_files = sorted(
                    name for name in package.namelist() if name.startswith("word/media/")
                )

            self.assertGreaterEqual(len(media_files), 2)

    def test_normalize_with_local_office_reports_backend_failure(self) -> None:
        processor = WordWatermarkProcessor()
        completed = subprocess.CompletedProcess(
            args=["powershell"],
            returncode=1,
            stdout="",
            stderr="Word.Application: 文件可能已经损坏。",
        )
        with tempfile.TemporaryDirectory() as tmp_dir:
            root = Path(tmp_dir)
            source = root / "input.docx"
            target = root / "target.docx"
            source.write_bytes(b"invalid")
            with mock.patch("ppt_watermark.word_processor.subprocess.run", return_value=completed):
                with self.assertRaisesRegex(ValueError, "Word/WPS 另存为标准"):
                    processor._normalize_with_local_office(source, target)


if __name__ == "__main__":
    unittest.main()
