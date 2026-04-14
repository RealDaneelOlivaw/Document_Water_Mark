import io
import tempfile
import unittest
import zipfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from ppt_watermark.config import WatermarkConfig
from ppt_watermark.ppt_processor import PptWatermarkProcessor


class PptProcessorTests(unittest.TestCase):
    def test_processor_generates_output_file(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            tmp_path = Path(tmp_dir)
            source_ppt = tmp_path / "input.pptx"
            output_ppt = tmp_path / "output.pptx"

            image_stream = io.BytesIO()
            Image.new("RGB", (320, 200), "navy").save(image_stream, format="PNG")
            image_stream.seek(0)

            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            slide.shapes.add_picture(image_stream, Inches(1), Inches(1), Inches(5), Inches(3))
            prs.save(str(source_ppt))

            processor = PptWatermarkProcessor()
            stats = processor.process_file(
                input_path=source_ppt,
                output_path=output_ppt,
                config=WatermarkConfig(text="CONFIDENTIAL", opacity=40),
            )

            self.assertTrue(output_ppt.exists())
            self.assertEqual(stats.total_images, 1)
            self.assertEqual(stats.watermarked_images, 1)

    def test_same_source_image_with_different_display_sizes_creates_distinct_media(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            tmp_path = Path(tmp_dir)
            source_ppt = tmp_path / "input.pptx"
            output_ppt = tmp_path / "output.pptx"

            image_path = tmp_path / "shared.png"
            Image.new("RGB", (600, 300), "white").save(image_path)

            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            slide.shapes.add_picture(str(image_path), Inches(1), Inches(1), Inches(2), Inches(1))
            slide.shapes.add_picture(str(image_path), Inches(1), Inches(3), Inches(4), Inches(2))
            prs.save(str(source_ppt))

            processor = PptWatermarkProcessor(min_area_pct=0)
            stats = processor.process_file(
                input_path=source_ppt,
                output_path=output_ppt,
                config=WatermarkConfig(text="CONFIDENTIAL", opacity=40),
            )

            self.assertEqual(stats.total_images, 2)
            self.assertEqual(stats.watermarked_images, 2)

            with zipfile.ZipFile(output_ppt) as package:
                media_files = sorted(
                    name for name in package.namelist() if name.startswith("ppt/media/")
                )

            self.assertGreaterEqual(len(media_files), 2)


if __name__ == "__main__":
    unittest.main()
